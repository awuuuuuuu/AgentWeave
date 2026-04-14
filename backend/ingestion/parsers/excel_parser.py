from __future__ import annotations

import logging
from pathlib import Path
from typing import Union

from openpyxl import load_workbook

from .base import BaseParser, ParsedChunk, register_parser
from .csv_parser import _detect_header, _extract_preamble, _row_to_text

logger = logging.getLogger(__name__)

@register_parser
class ExcelParser(BaseParser):
    """Excel（.xlsx）文件解析器

    核心思路：
    1. 每个 Sheet 独立处理，Sheet 名作为 section_path
    2. 每数据行 → 一个 chunk，格式：`字段: 值; 字段: 值`
    """

    supported_extensions = (".xlsx",)

    def parse(self, source: Union[str, Path, bytes]) -> list[ParsedChunk]:
        filename = self._source_name(source)
        raw_bytes = self._read_bytes(source)

        try:
            return _parse_excel(raw_bytes, filename)
        except Exception as exc:
            logger.warning("解析 Excel 文件 '%s' 失败: %s", filename, exc)
            return [ParsedChunk(text="", metadata={"source_file": filename, "content_type": "error", "section_path": "", "error": str(exc)})]

def _parse_excel(raw_bytes: bytes, filename: str) -> list[ParsedChunk]:
    import io

    wb = load_workbook(io.BytesIO(raw_bytes), read_only=False, data_only=True)
    try:
        chunks: list[ParsedChunk] = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            filled_map = _build_fill_map(sheet)
            rows = _sheet_to_rows(sheet, filled_map)
            sheet_chunks = _parse_sheet(rows, filename, sheet_name)
            chunks.extend(sheet_chunks)
            logger.debug("Sheet '%s'：提取 %d 个 chunk", sheet_name, len(sheet_chunks))
    finally:
        wb.close()
    return chunks

def _build_fill_map(sheet) -> dict[tuple[int, int], str]:
    """
    将合并的单元格拆分
    例：部门列合并了 10 行"销售部"，填充后每行都能看到"销售部"。
    """
    filled_map: dict[tuple[int, int], str] = {}
    for merge_range in sheet.merged_cells.ranges:
        top_left = sheet.cell(merge_range.min_row, merge_range.min_col).value
        value = "" if top_left is None else str(top_left).strip()
        for row in range(merge_range.min_row, merge_range.max_row + 1):
            for col in range(merge_range.min_col, merge_range.max_col + 1):
                if row == merge_range.min_row and col == merge_range.min_col:
                    continue
                filled_map[(row, col)] = value
    return filled_map

def _sheet_to_rows(sheet, filled_map: dict[tuple[int, int], str]) -> list[list[str]]:
    """将 sheet 所有行转为字符串二维列表"""
    rows = []
    for row in sheet.iter_rows():
        str_row = [_cell_value(cell, filled_map) for cell in row]
        rows.append(str_row)
    
    # 去掉末尾空行
    while rows and not any(rows[-1]):
        rows.pop()

    return rows

def _cell_value(cell, filled_map: dict[tuple[int, int], str]) -> str:
    """提取单元格值"""
    key = (cell.row, cell.column)
    if key in filled_map:
        return filled_map[key]
    
    val = cell.value
    if val is None:
        return ""
    return str(val).strip()

def _parse_sheet(
    rows: list[list[str]],
    filename: str,
    sheet_name: str
) -> list[ParsedChunk]:
    """将单个 Sheet 的行数据转为 ParsedChunk 列表"""
    if not rows:
        return []
    
    header_idx = _detect_header(rows)
    preamble = _extract_preamble(rows, header_idx) if header_idx > 0 else {}

    if header_idx == -1:
        headers = [f"col_{i}" for i in range(len(rows[0]))]
        data_rows = rows
    else:
        headers = [cell or f"col_{i}" for i, cell in enumerate(rows[header_idx])]
        data_rows = rows[header_idx + 1:]

    section_path = sheet_name
    chunks: list[ParsedChunk] = []

    for row_num, row in enumerate(data_rows, start=header_idx + 2):
        if not any(row):
            continue
        text = _row_to_text(headers, row)
        if text:
            chunks.append(ParsedChunk(
                text=text,
                metadata={
                    "source_file": filename,
                    "content_type": "table",
                    "section_path": section_path,
                    "sheet_name": sheet_name,
                    "row_number": row_num,
                    **preamble,
                },
            ))

    return chunks