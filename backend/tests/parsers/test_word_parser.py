"""
WordParser 测试套件

测试分层：
  单元测试（无外部依赖）：用 python-docx 程序化生成 .docx，直接测解析逻辑
    - 注册表
    - 标题识别（styleId 方式，兼容中文 Office）
    - 标题层级栈 + section_path
    - 表格：加粗表头 / 无加粗回退 / 全空首行 / 合并单元格
    - 不可见字符清理（\\xa0、\\u2002）
    - 容错（损坏文件）

运行：
  uv run pytest backend/tests/parsers/test_word_parser.py -v
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from ingestion.parsers.base import ParsedChunk, get_parser
from ingestion.parsers.word_parser import WordParser, _build_section_path, _table_to_text


# ── 工具函数 ────────────────────────────────────────────────────────────────────

def _to_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_heading(doc: Document, text: str, level: int) -> None:
    """添加标题段落，styleId 为 'Heading{level}'。"""
    doc.add_heading(text, level=level)


def _set_row_bold(row, values: list[str]) -> None:
    """将已有行的单元格内容设为加粗。"""
    for i, val in enumerate(values):
        cell = row.cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(val)
        run.bold = True


def _add_bold_row(table, values: list[str]) -> None:
    """向表格追加一行，所有单元格内容加粗。"""
    row = table.add_row()
    _set_row_bold(row, values)


# ── Fixtures ────────────────────────────────────────────────────────────────────

@pytest.fixture
def simple_docx() -> bytes:
    """基础文档：一个 Heading 1 + 两段正文。"""
    doc = Document()
    _add_heading(doc, "Introduction", level=1)
    doc.add_paragraph("First paragraph of content.")
    doc.add_paragraph("Second paragraph of content.")
    return _to_bytes(doc)


@pytest.fixture
def heading_hierarchy_docx() -> bytes:
    """多层标题文档，验证 section_path 层级栈逻辑。

    结构：
        H1: Chapter 1
            H2: Section 1.1
                正文 A
            H2: Section 1.2
                正文 B
        H1: Chapter 2
            正文 C
    """
    doc = Document()
    _add_heading(doc, "Chapter 1", level=1)
    _add_heading(doc, "Section 1.1", level=2)
    doc.add_paragraph("Body text A.")
    _add_heading(doc, "Section 1.2", level=2)
    doc.add_paragraph("Body text B.")
    _add_heading(doc, "Chapter 2", level=1)
    doc.add_paragraph("Body text C.")
    return _to_bytes(doc)


@pytest.fixture
def table_bold_header_docx() -> bytes:
    """含加粗表头的表格。

    表头行加粗 → _is_row_bold 返回 True → 生成键值对格式。
    """
    doc = Document()
    doc.add_paragraph("Product table:")
    table = doc.add_table(rows=1, cols=3)
    # 直接把第一行（index 0）设为加粗表头，不另追加行
    _set_row_bold(table.rows[0], ["Product", "Quantity", "Price"])
    row1 = table.add_row()
    row1.cells[0].text = "Widget A"
    row1.cells[1].text = "100"
    row1.cells[2].text = "$9.99"
    row2 = table.add_row()
    row2.cells[0].text = "Widget B"
    row2.cells[1].text = "250"
    row2.cells[2].text = "$4.99"
    return _to_bytes(doc)


@pytest.fixture
def table_no_bold_docx() -> bytes:
    """无加粗的表格（首行非加粗）→ 回退到竖线拼接。"""
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    data = [["Name", "Score"], ["Alice", "95"], ["Bob", "87"]]
    for i, row_data in enumerate(data):
        for j, val in enumerate(row_data):
            table.rows[i].cells[j].text = val
    return _to_bytes(doc)


@pytest.fixture
def table_section_path_docx() -> bytes:
    """标题下方紧跟表格，验证表格的 section_path 携带正确的标题路径。"""
    doc = Document()
    _add_heading(doc, "Financial Results", level=1)
    _add_heading(doc, "Q1 Summary", level=2)
    table = doc.add_table(rows=2, cols=2)
    _add_bold_row(table, ["Item", "Value"])
    table.rows[1].cells[0].text = "Revenue"
    table.rows[1].cells[1].text = "$1M"
    return _to_bytes(doc)


@pytest.fixture
def nbsp_docx() -> bytes:
    """含 \\xa0（non-breaking space）和 \\u2002 的段落，验证空白清理。"""
    doc = Document()
    _add_heading(doc, "Title", level=1)
    # python-docx add_paragraph 会直接存入文字，\xa0 不被 strip() 清除
    para = doc.add_paragraph()
    para.add_run("Hello\xa0World\u2002Test")
    return _to_bytes(doc)


@pytest.fixture
def corrupted_docx() -> bytes:
    """非法字节流，测试容错处理。"""
    return b"PK\x03\x04this is not a valid docx file"


@pytest.fixture
def empty_docx() -> bytes:
    """空文档（无段落），应返回空列表。"""
    doc = Document()
    return _to_bytes(doc)


# ── 基础结构测试 ────────────────────────────────────────────────────────────────

class TestWordParserBasic:

    def test_registered_for_docx_extension(self):
        parser = get_parser(".docx")
        assert isinstance(parser, WordParser)

    def test_parse_returns_list(self, simple_docx):
        result = WordParser().parse(simple_docx)
        assert isinstance(result, list)

    def test_parse_returns_parsed_chunks(self, simple_docx):
        result = WordParser().parse(simple_docx)
        assert len(result) > 0
        assert all(isinstance(c, ParsedChunk) for c in result)

    def test_chunk_has_required_metadata(self, simple_docx):
        chunks = WordParser().parse(simple_docx)
        for chunk in chunks:
            assert "source_file" in chunk.metadata
            assert "content_type" in chunk.metadata
            assert "section_path" in chunk.metadata

    def test_empty_docx_returns_empty_list(self, empty_docx):
        # 空文档（Word 默认有一个空段落）不应返回 chunk
        result = WordParser().parse(empty_docx)
        assert result == [] or all(c.text == "" for c in result)

    def test_parse_from_path(self, simple_docx, tmp_path):
        path = tmp_path / "test.docx"
        path.write_bytes(simple_docx)
        result = WordParser().parse(path)
        assert len(result) > 0

    def test_source_file_from_bytes_is_unknown(self, simple_docx):
        chunks = WordParser().parse(simple_docx)
        assert all(c.metadata["source_file"] == "unknown" for c in chunks)

    def test_source_file_from_path(self, simple_docx, tmp_path):
        path = tmp_path / "report.docx"
        path.write_bytes(simple_docx)
        chunks = WordParser().parse(path)
        assert all(c.metadata["source_file"] == "report.docx" for c in chunks)


# ── 标题识别测试 ────────────────────────────────────────────────────────────────

class TestHeadingDetection:

    def test_heading_content_type_is_title(self, simple_docx):
        chunks = WordParser().parse(simple_docx)
        heading_chunks = [c for c in chunks if c.text == "Introduction"]
        assert len(heading_chunks) == 1
        assert heading_chunks[0].metadata["content_type"] == "title"

    def test_body_content_type_is_text(self, simple_docx):
        chunks = WordParser().parse(simple_docx)
        body_chunks = [c for c in chunks if c.metadata["content_type"] == "text"]
        assert len(body_chunks) >= 2

    def test_heading_level_in_metadata(self, simple_docx):
        chunks = WordParser().parse(simple_docx)
        heading = next(c for c in chunks if c.text == "Introduction")
        assert heading.metadata["heading_level"] == 1

    def test_non_heading_has_no_heading_level(self, simple_docx):
        chunks = WordParser().parse(simple_docx)
        body = next(c for c in chunks if c.metadata["content_type"] == "text")
        assert "heading_level" not in body.metadata

    def test_multiple_heading_levels(self, heading_hierarchy_docx):
        chunks = WordParser().parse(heading_hierarchy_docx)
        h1_chunks = [c for c in chunks if c.metadata.get("heading_level") == 1]
        h2_chunks = [c for c in chunks if c.metadata.get("heading_level") == 2]
        assert len(h1_chunks) == 2
        assert len(h2_chunks) == 2

    def test_style_name_in_metadata(self, simple_docx):
        chunks = WordParser().parse(simple_docx)
        heading = next(c for c in chunks if c.text == "Introduction")
        assert "Heading" in heading.metadata["style"]


# ── section_path 测试 ───────────────────────────────────────────────────────────

class TestSectionPath:

    def test_body_under_h1_has_correct_path(self, heading_hierarchy_docx):
        chunks = WordParser().parse(heading_hierarchy_docx)
        body_a = next(c for c in chunks if c.text == "Body text A.")
        assert "Section 1.1" in body_a.metadata["section_path"]
        assert "Chapter 1" in body_a.metadata["section_path"]

    def test_h2_clears_previous_h2(self, heading_hierarchy_docx):
        """Section 1.2 应替换 Section 1.1，不应同时出现。"""
        chunks = WordParser().parse(heading_hierarchy_docx)
        body_b = next(c for c in chunks if c.text == "Body text B.")
        assert "Section 1.2" in body_b.metadata["section_path"]
        assert "Section 1.1" not in body_b.metadata["section_path"]

    def test_h1_clears_h2(self, heading_hierarchy_docx):
        """进入 Chapter 2 后，Section 1.x 应被清除。"""
        chunks = WordParser().parse(heading_hierarchy_docx)
        body_c = next(c for c in chunks if c.text == "Body text C.")
        assert "Chapter 2" in body_c.metadata["section_path"]
        assert "Section" not in body_c.metadata["section_path"]

    def test_section_path_separator(self, heading_hierarchy_docx):
        chunks = WordParser().parse(heading_hierarchy_docx)
        body_a = next(c for c in chunks if c.text == "Body text A.")
        assert " > " in body_a.metadata["section_path"]

    def test_heading_itself_has_section_path(self, heading_hierarchy_docx):
        """标题 chunk 的 section_path 应包含自身（入栈后读取）。"""
        chunks = WordParser().parse(heading_hierarchy_docx)
        h2 = next(c for c in chunks if c.text == "Section 1.1")
        assert "Section 1.1" in h2.metadata["section_path"]

    def test_build_section_path_empty(self):
        assert _build_section_path({}) == ""

    def test_build_section_path_single(self):
        assert _build_section_path({1: "Chapter 1"}) == "Chapter 1"

    def test_build_section_path_nested(self):
        result = _build_section_path({1: "Chapter 1", 2: "Section 1.1"})
        assert result == "Chapter 1 > Section 1.1"


# ── 表格测试 ────────────────────────────────────────────────────────────────────

class TestTableParsing:

    def test_table_content_type(self, table_bold_header_docx):
        chunks = WordParser().parse(table_bold_header_docx)
        table_chunks = [c for c in chunks if c.metadata["content_type"] == "table"]
        assert len(table_chunks) == 1

    def test_bold_header_generates_key_value(self, table_bold_header_docx):
        """加粗表头 → 'Product: Widget A; Quantity: 100; Price: $9.99' 格式。"""
        chunks = WordParser().parse(table_bold_header_docx)
        table_chunk = next(c for c in chunks if c.metadata["content_type"] == "table")
        assert "Product: Widget A" in table_chunk.text
        assert "Quantity: 100" in table_chunk.text
        assert "Price: $9.99" in table_chunk.text

    def test_no_bold_falls_back_to_pipe(self, table_no_bold_docx):
        """无加粗 → 竖线拼接回退，不生成错误键值对。"""
        chunks = WordParser().parse(table_no_bold_docx)
        table_chunk = next(c for c in chunks if c.metadata["content_type"] == "table")
        assert "|" in table_chunk.text

    def test_table_section_path(self, table_section_path_docx):
        """表格的 section_path 应包含上方的标题路径。"""
        chunks = WordParser().parse(table_section_path_docx)
        table_chunk = next(c for c in chunks if c.metadata["content_type"] == "table")
        assert "Financial Results" in table_chunk.metadata["section_path"]
        assert "Q1 Summary" in table_chunk.metadata["section_path"]

    def test_table_to_text_empty(self):
        """空表格返回空字符串。"""
        doc = Document()
        table = doc.add_table(rows=0, cols=0)
        assert _table_to_text(table) == ""

    def test_table_to_text_single_row(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=3)
        for i, val in enumerate(["A", "B", "C"]):
            table.rows[0].cells[i].text = val
        result = _table_to_text(table)
        assert "A" in result and "B" in result and "C" in result


# ── _table_to_text 函数级精确测试（建议 1）──────────────────────────────────────

class TestTableToTextDirect:
    """直接测 _table_to_text()，覆盖所有分支，不经过 parse() 间接触发。"""

    def _make_table(self, data: list[list[str]], bold_first_row: bool = False) -> object:
        """程序化构建 Table 对象。"""
        cols = max(len(row) for row in data) if data else 1
        doc = Document()
        table = doc.add_table(rows=len(data), cols=cols)
        for i, row_data in enumerate(data):
            for j, val in enumerate(row_data):
                cell = table.rows[i].cells[j]
                if bold_first_row and i == 0:
                    cell.text = ""
                    cell.paragraphs[0].add_run(val).bold = True
                else:
                    cell.text = val
        return table

    def test_kv_exact_output(self):
        """加粗表头 → 精确的键值对格式。"""
        table = self._make_table(
            [["Product", "Price"], ["Widget A", "$9.99"], ["Widget B", "$4.99"]],
            bold_first_row=True,
        )
        result = _table_to_text(table)
        assert result == "Product: Widget A; Price: $9.99\nProduct: Widget B; Price: $4.99"

    def test_no_bold_fallback_pipe(self):
        """无加粗 → 竖线拼接，不生成错误键值对。"""
        table = self._make_table([["Name", "Score"], ["Alice", "95"], ["Bob", "87"]])
        result = _table_to_text(table)
        assert "|" in result
        assert "Name: Alice" not in result  # 没有错误地当成表头

    def test_empty_header_cell_skipped(self):
        """表头某列为空时，该列的键值对被跳过（不生成 ': value'）。"""
        table = self._make_table(
            [["Product", "", "Price"], ["Widget A", "100", "$9.99"]],
            bold_first_row=True,
        )
        result = _table_to_text(table)
        assert "Product: Widget A" in result
        assert "Price: $9.99" in result
        assert ": 100" not in result  # 空表头对应的值被跳过

    def test_merged_cells_deduped(self):
        """合并单元格在 python-docx 中重复返回，seen 集合应去重。"""
        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        # 手动设置加粗表头
        for j, h in enumerate(["A", "B", "C"]):
            cell = table.rows[0].cells[j]
            cell.text = ""
            cell.paragraphs[0].add_run(h).bold = True
        # 合并数据行第 1、2 列
        table.rows[1].cells[0].text = "v1"
        table.rows[1].cells[1].merge(table.rows[1].cells[2])
        table.rows[1].cells[1].text = "v2"
        result = _table_to_text(table)
        # 合并列不应产生重复的 "B: v2; B: v2"
        assert result.count("B: v2") == 1


# ── 文档顺序稳定性测试（建议 2）──────────────────────────────────────────────────

class TestOutputOrder:
    """验证输出 chunk 顺序与文档中的物理顺序一致。"""

    def test_heading_before_body(self, simple_docx):
        chunks = WordParser().parse(simple_docx)
        texts = [c.text for c in chunks]
        intro_idx = texts.index("Introduction")
        first_para_idx = texts.index("First paragraph of content.")
        assert intro_idx < first_para_idx

    def test_heading_before_table(self, table_section_path_docx):
        chunks = WordParser().parse(table_section_path_docx)
        heading_idx = next(
            i for i, c in enumerate(chunks) if c.metadata["content_type"] == "title"
        )
        table_idx = next(
            i for i, c in enumerate(chunks) if c.metadata["content_type"] == "table"
        )
        assert heading_idx < table_idx

    def test_section_order_preserved(self, heading_hierarchy_docx):
        """多节文档：Chapter 1 的内容先于 Chapter 2 出现。"""
        chunks = WordParser().parse(heading_hierarchy_docx)
        texts = [c.text for c in chunks]
        assert texts.index("Body text A.") < texts.index("Body text C.")


# ── 标题跳级测试（建议 3）──────────────────────────────────────────────────────

class TestHeadingSkipLevel:

    def test_skip_level_section_path(self):
        """H1 → H3（跳过 H2），section_path 应包含两者，不含虚构的中间层。"""
        doc = Document()
        _add_heading(doc, "Chapter 1", level=1)
        _add_heading(doc, "Deep Section", level=3)
        doc.add_paragraph("Body text.")
        chunks = WordParser().parse(_to_bytes(doc))
        body = next(c for c in chunks if c.text == "Body text.")
        assert "Chapter 1" in body.metadata["section_path"]
        assert "Deep Section" in body.metadata["section_path"]

    def test_lower_heading_clears_higher_skip(self):
        """H1 → H3 → H2：H2 出现时应清除 H3，保留 H1。"""
        doc = Document()
        _add_heading(doc, "Chapter 1", level=1)
        _add_heading(doc, "Deep Section", level=3)
        _add_heading(doc, "Normal Section", level=2)
        doc.add_paragraph("After H2.")
        chunks = WordParser().parse(_to_bytes(doc))
        body = next(c for c in chunks if c.text == "After H2.")
        path = body.metadata["section_path"]
        assert "Chapter 1" in path
        assert "Normal Section" in path
        assert "Deep Section" not in path  # H3 被 H2 清除

    def test_build_section_path_missing_levels(self):
        """_build_section_path 直接接收不连续层级，输出应只含有值的层级。"""
        result = _build_section_path({1: "Chapter 1", 3: "Deep Section"})
        assert result == "Chapter 1 > Deep Section"

    def test_build_section_path_only_deep(self):
        """文档从 H3 开始（无 H1/H2），section_path 只输出 H3 文字。"""
        result = _build_section_path({3: "Orphan Section"})
        assert result == "Orphan Section"


# ── 列表样式测试（建议 5）──────────────────────────────────────────────────────

class TestListParagraphs:

    def test_list_item_extracted_as_text(self):
        """列表项（List Bullet / List Number）应被提取为 content_type='text'。"""
        doc = Document()
        _add_heading(doc, "My List", level=1)
        doc.add_paragraph("Item one", style="List Bullet")
        doc.add_paragraph("Item two", style="List Number")
        chunks = WordParser().parse(_to_bytes(doc))
        list_chunks = [c for c in chunks if c.metadata["content_type"] == "text"]
        texts = [c.text for c in list_chunks]
        assert "Item one" in texts
        assert "Item two" in texts

    def test_list_item_not_treated_as_title(self):
        """列表项不应被识别为 title。"""
        doc = Document()
        doc.add_paragraph("Bullet point", style="List Bullet")
        chunks = WordParser().parse(_to_bytes(doc))
        assert all(c.metadata["content_type"] != "title" for c in chunks)


# ── metadata schema 一致性测试（建议 7）─────────────────────────────────────────

class TestMetadataSchema:

    # text/title chunk 必须包含的字段
    _TEXT_KEYS = {"source_file", "content_type", "style", "section_path"}
    # table chunk 必须包含的字段
    _TABLE_KEYS = {"source_file", "content_type", "section_path"}

    def test_text_chunks_have_consistent_schema(self, heading_hierarchy_docx):
        chunks = WordParser().parse(heading_hierarchy_docx)
        text_chunks = [c for c in chunks if c.metadata["content_type"] in {"text", "title"}]
        for chunk in text_chunks:
            missing = self._TEXT_KEYS - chunk.metadata.keys()
            assert not missing, f"chunk '{chunk.text[:30]}' 缺少字段: {missing}"

    def test_table_chunks_have_consistent_schema(self, table_section_path_docx):
        chunks = WordParser().parse(table_section_path_docx)
        table_chunks = [c for c in chunks if c.metadata["content_type"] == "table"]
        for chunk in table_chunks:
            missing = self._TABLE_KEYS - chunk.metadata.keys()
            assert not missing, f"table chunk 缺少字段: {missing}"

    def test_title_chunks_have_heading_level(self, heading_hierarchy_docx):
        chunks = WordParser().parse(heading_hierarchy_docx)
        title_chunks = [c for c in chunks if c.metadata["content_type"] == "title"]
        for chunk in title_chunks:
            assert "heading_level" in chunk.metadata
            assert isinstance(chunk.metadata["heading_level"], int)

    def test_content_type_values_valid(self, heading_hierarchy_docx):
        valid = {"text", "title", "table"}
        chunks = WordParser().parse(heading_hierarchy_docx)
        for chunk in chunks:
            assert chunk.metadata["content_type"] in valid


# ── 不可见字符清理测试 ──────────────────────────────────────────────────────────

class TestWhitespaceCleaning:

    def test_nbsp_is_normalized(self, nbsp_docx):
        """\\xa0 应被规范化为普通空格。"""
        chunks = WordParser().parse(nbsp_docx)
        body = next(c for c in chunks if c.metadata["content_type"] == "text")
        assert "\xa0" not in body.text
        assert "Hello World Test" in body.text

    def test_no_double_spaces(self, nbsp_docx):
        chunks = WordParser().parse(nbsp_docx)
        for chunk in chunks:
            assert "  " not in chunk.text  # 不含连续空格


# ── 容错测试 ────────────────────────────────────────────────────────────────────

class TestErrorHandling:

    def test_corrupted_docx_no_exception(self, corrupted_docx):
        result = WordParser().parse(corrupted_docx)
        assert isinstance(result, list)

    def test_corrupted_docx_has_error_metadata(self, corrupted_docx):
        result = WordParser().parse(corrupted_docx)
        assert len(result) == 1
        assert "error" in result[0].metadata
        assert result[0].text == ""
