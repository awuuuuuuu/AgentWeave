"""
ExcelParser 测试套件

运行：
  uv run pytest backend/tests/parsers/test_excel_parser.py -v
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from openpyxl import Workbook

from ingestion.parsers.base import ParsedChunk, get_parser
from ingestion.parsers.excel_parser import ExcelParser


# ── Fixture 工厂 ─────────────────────────────────────────────────────────────────

def _make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    """程序化构造 xlsx 字节流，无需外部文件。

    sheets: {sheet_name: [[row1_col1, row1_col2, ...], [row2_...], ...]}
    """
    wb = Workbook()
    wb.remove(wb.active)  # 删除默认空 Sheet
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse(sheets: dict[str, list[list]]) -> list[ParsedChunk]:
    return ExcelParser().parse(_make_xlsx(sheets))


# ── 基础结构 ────────────────────────────────────────────────────────────────────

class TestExcelParserBasic:

    def test_registered_for_xlsx(self):
        assert isinstance(get_parser(".xlsx"), ExcelParser)

    def test_returns_list(self):
        result = _parse({"Sheet1": [["name", "age"], ["Alice", 30]]})
        assert isinstance(result, list)

    def test_empty_sheet_returns_empty(self):
        assert _parse({"Sheet1": []}) == []

    def test_chunk_has_required_metadata(self):
        chunks = _parse({"Sheet1": [["name", "age"], ["Alice", 30]]})
        for c in chunks:
            assert "source_file" in c.metadata
            assert "content_type" in c.metadata
            assert "section_path" in c.metadata
            assert "sheet_name" in c.metadata
            assert "row_number" in c.metadata

    def test_content_type_is_table(self):
        chunks = _parse({"Sheet1": [["name", "age"], ["Alice", 30]]})
        assert all(c.metadata["content_type"] == "table" for c in chunks)

    def test_source_file_from_path(self, tmp_path):
        p = tmp_path / "data.xlsx"
        p.write_bytes(_make_xlsx({"Sheet1": [["name"], ["Alice"]]}))
        chunks = ExcelParser().parse(p)
        assert chunks[0].metadata["source_file"] == "data.xlsx"


# ── Sheet 处理 ──────────────────────────────────────────────────────────────────

class TestSheetHandling:

    def test_sheet_name_is_section_path(self):
        chunks = _parse({"财务报表": [["项目", "金额"], ["收入", 100]]})
        assert chunks[0].metadata["section_path"] == "财务报表"
        assert chunks[0].metadata["sheet_name"] == "财务报表"

    def test_multiple_sheets_all_parsed(self):
        chunks = _parse({
            "Sheet1": [["name"], ["Alice"], ["Bob"]],
            "Sheet2": [["product"], ["Widget"], ["Gadget"]],
        })
        sheet_names = {c.metadata["sheet_name"] for c in chunks}
        assert "Sheet1" in sheet_names
        assert "Sheet2" in sheet_names

    def test_multiple_sheets_chunk_count(self):
        chunks = _parse({
            "A": [["name", "age"], ["Alice", 30], ["Bob", 25]],
            "B": [["product", "price"], ["Widget", 9.99]],
        })
        assert len(chunks) == 3  # 2 + 1

    def test_empty_sheet_skipped(self):
        chunks = _parse({
            "Empty": [],
            "Data": [["name"], ["Alice"]],
        })
        sheet_names = {c.metadata["sheet_name"] for c in chunks}
        assert "Empty" not in sheet_names
        assert "Data" in sheet_names


# ── 数据解析 ────────────────────────────────────────────────────────────────────

class TestExcelDataParsing:

    def test_one_row_one_chunk(self):
        chunks = _parse({"Sheet1": [["name", "age"], ["Alice", 30], ["Bob", 25]]})
        assert len(chunks) == 2

    def test_chunk_text_format(self):
        chunks = _parse({"Sheet1": [["name", "age"], ["Alice", 30]]})
        assert "name: Alice" in chunks[0].text
        assert "age: 30" in chunks[0].text

    def test_blank_rows_skipped(self):
        chunks = _parse({"Sheet1": [["name", "age"], ["Alice", 30], [None, None], ["Bob", 25]]})
        assert len(chunks) == 2

    def test_row_number_in_metadata(self):
        chunks = _parse({"Sheet1": [["name", "age"], ["Alice", 30], ["Bob", 25]]})
        assert chunks[0].metadata["row_number"] == 2
        assert chunks[1].metadata["row_number"] == 3

    def test_none_values_treated_as_empty(self):
        chunks = _parse({"Sheet1": [["name", "age"], ["Alice", None]]})
        assert "age" not in chunks[0].text
        assert "name: Alice" in chunks[0].text

    def test_numeric_values_converted_to_string(self):
        chunks = _parse({"Sheet1": [["product", "price"], ["Widget", 9.99]]})
        assert "9.99" in chunks[0].text

    def test_trailing_empty_rows_ignored(self):
        chunks = _parse({"Sheet1": [["name", "age"], ["Alice", 30], [None, None], [None, None]]})
        assert len(chunks) == 1


# ── 表头检测 ────────────────────────────────────────────────────────────────────

class TestExcelHeaderDetection:

    def test_first_row_as_header(self):
        chunks = _parse({"Sheet1": [["name", "age"], ["Alice", 30]]})
        assert "name: Alice" in chunks[0].text

    def test_leading_empty_row_skipped_for_header(self):
        chunks = _parse({"Sheet1": [[None, None], ["name", "age"], ["Alice", 30]]})
        assert "name: Alice" in chunks[0].text

    def test_single_column_falls_back_to_col0(self):
        # 单列无法区分表头与数据，全部作为数据行，用 col_0 命名
        chunks = _parse({"Sheet1": [["Alice"], ["Bob"], ["Carol"]]})
        assert all("col_0" in c.text for c in chunks)


# ── 容错 ────────────────────────────────────────────────────────────────────────

class TestPreamble:

    def test_preamble_rows_in_metadata(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Company: XX Group"])       # 单列冒号格式 preamble
        ws.append(["Period: 2024Q1"])
        ws.append(["name", "amount"])          # 表头（≥2 列，触发检测）
        ws.append(["Alice", 100])
        buf = io.BytesIO()
        wb.save(buf)
        chunks = ExcelParser().parse(buf.getvalue())
        assert chunks[0].metadata.get("Company") == "XX Group"
        assert chunks[0].metadata.get("Period") == "2024Q1"

    def test_preamble_propagated_to_all_chunks(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Source: system export"])   # 单列冒号格式 preamble
        ws.append(["name", "amount"])
        ws.append(["Alice", 100])
        ws.append(["Bob", 200])
        buf = io.BytesIO()
        wb.save(buf)
        chunks = ExcelParser().parse(buf.getvalue())
        assert all(c.metadata.get("Source") == "system export" for c in chunks)

    def test_no_preamble_no_extra_metadata(self):
        chunks = _parse({"Sheet1": [["name", "age"], ["Alice", 30]]})
        assert "row_0" not in chunks[0].metadata


class TestMergedCells:

    def test_merged_column_value_propagated(self):
        """合并单元格的值应填充到每一行，不能只出现在第一行。"""
        wb = Workbook()
        ws = wb.active
        ws.append(["department", "name"])
        ws.append(["Sales", "Alice"])
        ws.append(["Sales", "Bob"])     # 这两行的 department 列将被合并
        ws.append(["Sales", "Carol"])
        # 合并 A2:A4（Sales 占三行）
        ws.merge_cells("A2:A4")
        buf = io.BytesIO()
        wb.save(buf)

        chunks = ExcelParser().parse(buf.getvalue())
        # 三行数据都应该能看到 "Sales"
        texts = [c.text for c in chunks]
        assert all("department: Sales" in t for t in texts), (
            f"合并单元格的值未正确填充，实际输出：{texts}"
        )

    def test_merged_header_row_detected(self):
        """合并单元格出现在表头行时不应崩溃。"""
        wb = Workbook()
        ws = wb.active
        ws.append(["name", "score", "score"])  # 双列 score 表头
        ws.merge_cells("B1:C1")
        ws.append(["Alice", 90, 85])
        buf = io.BytesIO()
        wb.save(buf)
        chunks = ExcelParser().parse(buf.getvalue())
        assert isinstance(chunks, list)


class TestExcelErrorHandling:

    def test_invalid_bytes_returns_error_chunk(self):
        chunks = ExcelParser().parse(b"not an xlsx file")
        assert isinstance(chunks, list)
        assert len(chunks) == 1
        assert "error" in chunks[0].metadata
